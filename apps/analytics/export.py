"""
Выгрузка Аналитики 2.0 в Excel: любой список «под цифрой» и любая таблица
вкладки — одной функцией.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Iterable, List, Sequence

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E9EEF5")
_MONEY_FORMAT = "#,##0.00"
_DATE_FORMAT = "DD.MM.YYYY"
_MAX_WIDTH = 50


def build_workbook(
    title: str, headers: Sequence[str], rows: Iterable[Sequence], subtitle: str = ""
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]

    sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    header_row = 2
    if subtitle:
        sheet.cell(row=2, column=1, value=subtitle)
        header_row = 4

    widths = [len(str(header)) for header in headers]
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=header_row, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_index, row in enumerate(rows, header_row + 1):
        for col, value in enumerate(row, 1):
            cell = sheet.cell(row=row_index, column=col, value=value)
            if isinstance(value, Decimal):
                cell.number_format = _MONEY_FORMAT
            elif isinstance(value, date):
                cell.number_format = _DATE_FORMAT
            widths[col - 1] = max(widths[col - 1], len(_display(value)))

    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = min(
            width + 2, _MAX_WIDTH
        )
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    return workbook


def _display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return "00.00.0000"
    if isinstance(value, Decimal):
        return f"{value:,.2f}"
    return str(value)


def xlsx_response(workbook: Workbook, filename: str) -> HttpResponse:
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def table_response(
    title: str,
    headers: Sequence[str],
    rows: List[Sequence],
    filename: str,
    subtitle: str = "",
) -> HttpResponse:
    return xlsx_response(build_workbook(title, headers, rows, subtitle), filename)
