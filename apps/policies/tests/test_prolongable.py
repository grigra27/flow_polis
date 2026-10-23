"""Грузы не продлевают никогда: виды с needs_prolongation=False исключены
из раздела и экспорта «Пролонгация» и из метрик продлений."""

import importlib
import io
from datetime import timedelta

import pytest
from django.apps import apps as django_apps
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.analytics import portfolio
from apps.billing import prolongation_services as prol
from apps.clients.models import Client
from apps.core.services.dashboard_v2_service import DashboardV2Service
from apps.insurers.models import Branch, InsuranceType, Insurer
from apps.policies.models import Policy, prolongation_due_q


def _policy(number, insurance_type, end_date, start_date=None):
    return Policy.objects.create(
        policy_number=number,
        dfa_number=f"DFA-{number}",
        client=Client.objects.create(client_name=f"ЛП {number}"),
        insurer=Insurer.objects.get_or_create(insurer_name="СК")[0],
        branch=Branch.objects.get_or_create(branch_name="Филиал")[0],
        insurance_type=insurance_type,
        property_description="Объект",
        start_date=start_date or end_date - timedelta(days=365),
        end_date=end_date,
    )


@pytest.fixture
def end():
    # 15-е число следующего месяца: полисы в силе на начало месяца
    # и попадают в горизонт 90 дней дашборда.
    today = timezone.localdate()
    return (today.replace(day=1) + timedelta(days=32)).replace(day=15)


@pytest.fixture
def pair(db, end):
    kasko = InsuranceType.objects.create(name="КАСКО (тест)")
    cargo = InsuranceType.objects.create(name="Грузы (тест)", needs_prolongation=False)
    return {
        "kasko": _policy("KASKO-PROL", kasko, end),
        "cargo": _policy("CARGO-PROL", cargo, end),
    }


def test_new_insurance_type_needs_prolongation_by_default(db):
    assert InsuranceType.objects.create(name="Новый вид").needs_prolongation


def test_migration_marks_cargo(db):
    cargo = InsuranceType.objects.create(name="Грузы", needs_prolongation=True)
    migration = importlib.import_module(
        "apps.insurers.migrations.0015_insurancetype_needs_prolongation"
    )
    migration.mark_cargo(django_apps, None)
    cargo.refresh_from_db()
    assert not cargo.needs_prolongation


def test_queryset_prolongable(pair):
    assert list(Policy.objects.prolongable()) == [pair["kasko"]]


def test_prolongation_page_and_month_counts(pair, end):
    assert list(prol.get_prolongation_policies(end.year, end.month)) == [pair["kasko"]]

    months = [end.replace(day=1)]
    options = prol.build_prolongation_month_options(months, end.year, end.month)
    assert options[0]["total"] == 1


def test_prolongation_export(client, django_user_model, pair, end):
    user = django_user_model.objects.create_user("staff", password="x", is_staff=True)
    client.force_login(user)
    response = client.get(
        reverse("reports:export_policy_expiration"),
        {
            "date_from": end.replace(day=1).isoformat(),
            "date_to": (end + timedelta(days=1)).isoformat(),
        },
    )
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    text = " ".join(
        str(cell)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
    )
    assert "KASKO-PROL" in text
    assert "CARGO-PROL" not in text


def test_analytics_renewal_due(pair, end):
    due = portfolio.renewal_due(Policy.objects.all(), end.replace(day=1), end)
    assert list(due) == [pair["kasko"]]


def test_dashboard_renewal_radar(pair):
    radar = DashboardV2Service()._build_renewal_radar(
        active_policies_qs=Policy.objects.all(), today=timezone.localdate()
    )
    assert radar["active_total"] == 2
    assert radar["horizons"][-1]["count"] == 1


@pytest.fixture
def rule_a_cases(db, end):
    """Правило пролонгации «вариант А» (2026-10-23): полис доживает до своей
    даты окончания — активен и не расторгнут досрочно."""
    kasko = InsuranceType.objects.create(name="КАСКО (правило А)")

    def make(number, **fields):
        policy = _policy(number, kasko, end, start_date=fields.pop("start_date", None))
        Policy.objects.filter(pk=policy.pk).update(**fields)
        return number

    return {
        "in": {
            make("A-PLAIN"),
            # Расторжение не раньше окончания — не досрочное.
            make("A-TERM-AT-END", termination_date=end),
            # Начался в том же месяце — всё равно подлежит пролонгации.
            make("A-SHORT", start_date=end - timedelta(days=5)),
        },
        "out": {
            # Как №205: расторгнут за неделю до окончания, флаг снят.
            make(
                "A-TERM-EARLY",
                termination_date=end - timedelta(days=7),
                policy_active=False,
            ),
            # Как №891: дата расторжения есть, а флаг забыли снять.
            make("A-TERM-FLAG-LEFT", termination_date=end - timedelta(days=7)),
            make("A-INACTIVE", policy_active=False),
        },
    }


def test_prolongation_due_q(rule_a_cases):
    numbers = set(
        Policy.objects.filter(prolongation_due_q()).values_list(
            "policy_number", flat=True
        )
    )
    assert numbers == rule_a_cases["in"]


def test_page_and_export_use_the_same_rule(
    client, django_user_model, rule_a_cases, end
):
    page = set(
        prol.get_prolongation_policies(end.year, end.month).values_list(
            "policy_number", flat=True
        )
    )
    assert page == rule_a_cases["in"]

    user = django_user_model.objects.create_user("staff", password="x", is_staff=True)
    client.force_login(user)
    start, finish = prol.month_bounds(end.year, end.month)
    response = client.get(
        reverse("reports:export_policy_expiration"),
        {"date_from": start.isoformat(), "date_to": finish.isoformat()},
    )
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    text = " ".join(
        str(cell)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
    )
    for number in rule_a_cases["in"]:
        assert number in text
    for number in rule_a_cases["out"]:
        assert number not in text


def _cells(content):
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    return [
        list(row)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
    ]


def test_email_attachment_equals_export_for_the_month(
    client, django_user_model, rule_a_cases, end
):
    """Excel к письму со вкладки и экспорт за тот же месяц — одна таблица."""
    batch = prol.get_or_create_batch(end.year, end.month)
    attachment = prol.build_prolongation_attachment(
        batch, prol.get_prolongation_policies(end.year, end.month)
    )

    user = django_user_model.objects.create_user("staff", password="x", is_staff=True)
    client.force_login(user)
    response = client.get(
        reverse("reports:export_policy_expiration"),
        {
            "date_from": batch.starts_on.isoformat(),
            "date_to": batch.ends_on.isoformat(),
        },
    )

    assert _cells(attachment.read()) == _cells(response.content)


def test_faq_modals_on_both_pages(client, django_user_model, db):
    user = django_user_model.objects.create_user("staff", password="x", is_staff=True)
    client.force_login(user)
    rule = "Дата расторжения» пуста или не раньше даты окончания"

    page = client.get(reverse("policies:prolongation")).content.decode()
    assert 'id="prolongationFaqModal"' in page
    assert rule in page

    reports = client.get(reverse("reports:index")).content.decode()
    assert 'id="policyExpirationFaqModal"' in reports
    assert rule in reports
