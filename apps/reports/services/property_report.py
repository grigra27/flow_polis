"""
«Подробный отчёт по имуществу» для лизинговой компании.

Где и как застраховано её имущество на дату среза: страховщики, структура
имущества, филиалы и менеджеры, окончание покрытия, франшиза, крупнейшие
объекты. Основа — insured_value.insured_portfolio (полисы в силе с активным
ДФА, без ОСАГО, страховая сумма текущего года).

Решения владельца 2026-10-17: в отчёте нет КВ, премии и тарифа; блок
«возможно, не застраховано» во внешний отчёт не выносится (он во внутренней
аналитике, вкладка «Качество данных»).
"""

import calendar
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from typing import Any, Callable, Dict, List, Optional

from django.utils import timezone

from apps.analytics.money import HISTORY_RELIABLE_FROM, MONTH_SHORT
from apps.insurers.models import Insurer

from .insured_value import ZERO, InsuredObject, insured_portfolio
from .property_text import (
    CONDITION_NEW,
    CONDITION_UNKNOWN,
    CONDITION_USED,
    NOT_RECOGNIZED,
    detect_brand,
    detect_condition,
    property_age,
)

ENDING_SOON_DAYS = 90
TOP_OBJECTS = 20
TOP_BRANDS = 15
TOP_MANAGERS = 15
MATRIX_INSURERS = 8
DYNAMICS_QUARTERS = 4

SIZE_BUCKETS = [
    ("до 1 млн", 0, 1_000_000),
    ("1–3 млн", 1_000_000, 3_000_000),
    ("3–5 млн", 3_000_000, 5_000_000),
    ("5–10 млн", 5_000_000, 10_000_000),
    ("10–20 млн", 10_000_000, 20_000_000),
    ("20–50 млн", 20_000_000, 50_000_000),
    ("50 млн и больше", 50_000_000, None),
]
AGE_BUCKETS = [
    ("до 1 года", 0, 1),
    ("2–3 года", 2, 3),
    ("4–5 лет", 4, 5),
    ("6–10 лет", 6, 10),
    ("старше 10 лет", 11, None),
]
FRANCHISE_BUCKETS = [
    ("без франшизы", None, 0),
    ("до 30 тыс.", 0, 30_000),
    ("30–100 тыс.", 30_000, 100_000),
    ("больше 100 тыс.", 100_000, None),
]


def _share(part, total):
    return part / total * 100 if total else ZERO


def _group(
    objects: List[InsuredObject], key: Callable, label: Callable = None, order=None
) -> List[Dict[str, Any]]:
    """Группы: страховая стоимость, объектов, средняя, доля; сортировка по стоимости."""
    groups: Dict[Any, Dict[str, Any]] = {}
    for obj in objects:
        k = key(obj)
        group = groups.setdefault(
            k,
            {"key": k, "name": label(obj) if label else k, "sum": ZERO, "count": 0},
        )
        group["sum"] += obj.insured_sum
        group["count"] += 1
    total = sum((g["sum"] for g in groups.values()), ZERO)
    rows = list(groups.values())
    for row in rows:
        row["share"] = _share(row["sum"], total)
        row["avg"] = row["sum"] / row["count"] if row["count"] else ZERO
    if order:
        rows.sort(
            key=lambda r: order.index(r["key"]) if r["key"] in order else len(order)
        )
    else:
        rows.sort(key=lambda r: -r["sum"])
    _add_bars(rows)
    return rows


def _add_bars(rows, field="sum"):
    """Ширина полосы в % от максимума группы — строкой с точкой (для CSS)."""
    top = max((r[field] for r in rows), default=ZERO)
    for row in rows:
        row["bar"] = f"{float(row[field] / top * 100) if top else 0:.1f}"


def _size_bucket(value: Decimal) -> str:
    for label, low, high in SIZE_BUCKETS:
        if value >= low and (high is None or value < high):
            return label
    return SIZE_BUCKETS[-1][0]


def _age_bucket(age: Optional[int]) -> str:
    if age is None:
        return "год выпуска не указан"
    for label, low, high in AGE_BUCKETS:
        if age >= low and (high is None or age <= high):
            return label
    return AGE_BUCKETS[-1][0]


def _franchise_bucket(value: Decimal) -> str:
    if not value:
        return FRANCHISE_BUCKETS[0][0]
    for label, low, high in FRANCHISE_BUCKETS[1:]:
        if value > low and (high is None or value <= high):
            return label
    return FRANCHISE_BUCKETS[-1][0]


def _add_months(value: date, months: int) -> date:
    index = value.year * 12 + value.month - 1 + months
    year, month = divmod(index, 12)
    month += 1
    return date(year, month, min(value.day, calendar.monthrange(year, month)[1]))


def _quarter_ends_before(as_of: date, count: int) -> List[date]:
    first_month = (as_of.month - 1) // 3 * 3 + 1
    end = date(as_of.year, first_month, 1) - timedelta(days=1)
    result = []
    for _ in range(count):
        result.append(end)
        end = date(end.year, (end.month - 1) // 3 * 3 + 1, 1) - timedelta(days=1)
    return list(reversed(result))


def _totals(objects):
    total = sum((o.insured_sum for o in objects), ZERO)
    return {"sum": total, "count": len(objects)}


def _change(as_of, months, branch_ids, current):
    then = _add_months(as_of, -months)
    if then < HISTORY_RELIABLE_FROM:
        return None
    previous = _totals(insured_portfolio(then, branch_ids))
    return {
        "date": then,
        "sum": previous["sum"],
        "delta": current["sum"] - previous["sum"],
        "percent": _share(current["sum"] - previous["sum"], previous["sum"])
        if previous["sum"]
        else None,
    }


def _insurers(objects, as_of):
    ending_limit = as_of + timedelta(days=ENDING_SOON_DAYS)
    rows = _group(objects, key=lambda o: o.insurer_id, label=lambda o: o.insurer)
    ending = defaultdict(lambda: ZERO)
    for obj in objects:
        if obj.end_date <= ending_limit:
            ending[obj.insurer_id] += obj.insured_sum
    logos = {
        insurer.id: insurer.logo.url
        for insurer in Insurer.objects.filter(id__in=[r["key"] for r in rows]).only(
            "id", "logo"
        )
        if insurer.logo
    }
    for row in rows:
        row["ending_soon"] = ending[row["key"]]
        row["logo_url"] = logos.get(row["key"], "")
    top3 = sum((r["share"] for r in rows[:3]), ZERO)
    hhi = sum((r["share"] ** 2 for r in rows), ZERO)
    return rows, {"top3_share": top3, "hhi": hhi, "insurers": len(rows)}


def _matrix(objects, row_key, row_label, col_key, col_label, columns_limit=None):
    """Матрица страховой стоимости: строки × столбцы (столбцы по убыванию)."""
    col_totals = defaultdict(lambda: ZERO)
    col_names = {}
    cells = defaultdict(lambda: ZERO)
    row_totals = defaultdict(lambda: ZERO)
    row_names = {}
    for obj in objects:
        r, c = row_key(obj), col_key(obj)
        cells[(r, c)] += obj.insured_sum
        col_totals[c] += obj.insured_sum
        row_totals[r] += obj.insured_sum
        col_names[c] = col_label(obj)
        row_names[r] = row_label(obj)
    columns = sorted(col_totals, key=lambda c: -col_totals[c])
    other = []
    if columns_limit and len(columns) > columns_limit:
        columns, other = columns[:columns_limit], columns[columns_limit:]
    rows = []
    for r in sorted(row_totals, key=lambda r: -row_totals[r]):
        values = [cells[(r, c)] for c in columns]
        if other:
            values.append(sum((cells[(r, c)] for c in other), ZERO))
        rows.append(
            {
                "name": row_names[r],
                "total": row_totals[r],
                "cells": [
                    {
                        "sum": v,
                        "share": _share(v, row_totals[r]),
                        # Интенсивность заливки ячейки: доля в строке, 0..0.55.
                        "alpha": f"{float(_share(v, row_totals[r])) / 100 * 0.55:.3f}",
                    }
                    for v in values
                ],
            }
        )
    headers = [col_names[c] for c in columns] + (["Прочие"] if other else [])
    return {"headers": headers, "rows": rows}


def _insurer_dynamics(as_of, branch_ids, current_rows):
    dates = [
        d
        for d in _quarter_ends_before(as_of, DYNAMICS_QUARTERS)
        if d >= HISTORY_RELIABLE_FROM
    ]
    if not dates:
        return None
    snapshots = []
    for snapshot_date in dates:
        rows = _group(
            insured_portfolio(snapshot_date, branch_ids),
            key=lambda o: o.insurer_id,
            label=lambda o: o.insurer,
        )
        snapshots.append((snapshot_date, {r["key"]: r["share"] for r in rows}))
    snapshots.append((as_of, {r["key"]: r["share"] for r in current_rows}))
    names = {r["key"]: r["name"] for r in current_rows}
    top = [r["key"] for r in current_rows[:MATRIX_INSURERS]]
    return {
        "dates": [d for d, _ in snapshots],
        "rows": [
            {
                "name": names[pk],
                "shares": [shares.get(pk, ZERO) for _, shares in snapshots],
                "change": snapshots[-1][1].get(pk, ZERO)
                - snapshots[0][1].get(pk, ZERO),
            }
            for pk in top
        ],
    }


def _ending_by_month(objects, as_of, months=12):
    start = as_of.replace(day=1)
    buckets = []
    for offset in range(months):
        month = _add_months(start, offset)
        last = date(
            month.year, month.month, calendar.monthrange(month.year, month.month)[1]
        )
        selected = [o for o in objects if max(month, as_of) <= o.end_date <= last]
        buckets.append(
            {
                "label": f"{MONTH_SHORT[month.month - 1]} {month:%y}",
                "count": len(selected),
                "sum": sum((o.insured_sum for o in selected), ZERO),
            }
        )
    _add_bars(buckets)
    return buckets


def build_property_report(
    as_of: Optional[date] = None, branch_ids=None
) -> Dict[str, Any]:
    as_of = as_of or timezone.localdate()
    objects = insured_portfolio(as_of, branch_ids)
    totals = _totals(objects)
    brands = {o.id: detect_brand(o.description) for o in objects}
    conditions = {o.id: detect_condition(o.description) for o in objects}

    insurer_rows, concentration = _insurers(objects, as_of)

    brand_rows = _group(objects, key=lambda o: brands[o.id])
    recognized = [r for r in brand_rows if r["key"] != NOT_RECOGNIZED]
    not_recognized = next((r for r in brand_rows if r["key"] == NOT_RECOGNIZED), None)
    brand_top = recognized[:TOP_BRANDS]
    brand_other = recognized[TOP_BRANDS:]
    if brand_other:
        other_sum = sum((r["sum"] for r in brand_other), ZERO)
        brand_top.append(
            {
                "key": "other",
                "name": f"Другие марки ({len(brand_other)})",
                "sum": other_sum,
                "count": sum(r["count"] for r in brand_other),
                "share": _share(other_sum, totals["sum"]),
                "avg": ZERO,
            }
        )
    if not_recognized:
        brand_top.append(not_recognized)
    _add_bars(brand_top)

    managers = _group(
        [o for o in objects if o.manager_id],
        key=lambda o: o.manager_id,
        label=lambda o: f"{o.manager} ({o.branch})",
    )
    without_manager = [o for o in objects if not o.manager_id]

    end_years = _group(objects, key=lambda o: o.end_date.year)
    end_years.sort(key=lambda r: r["key"])
    for row in end_years:
        row["name"] = str(row["key"])

    top_objects = sorted(objects, key=lambda o: -o.insured_sum)[:TOP_OBJECTS]

    return {
        "as_of": as_of,
        "branch_ids": branch_ids or [],
        "totals": {
            **totals,
            "avg": totals["sum"] / totals["count"] if totals["count"] else ZERO,
            "lessees": len({o.client_id for o in objects}),
            "insurers": concentration["insurers"],
            "change_3m": _change(as_of, 3, branch_ids, totals),
            "change_12m": _change(as_of, 12, branch_ids, totals),
        },
        "insurers": insurer_rows,
        "concentration": concentration,
        "insurer_by_type": _matrix(
            objects,
            lambda o: o.insurer_id,
            lambda o: o.insurer,
            lambda o: o.type_id,
            lambda o: o.type,
        ),
        "insurer_dynamics": _insurer_dynamics(as_of, branch_ids, insurer_rows),
        "by_type": _group(objects, key=lambda o: o.type_id, label=lambda o: o.type),
        "by_size": _group(
            objects,
            key=lambda o: _size_bucket(o.insured_sum),
            order=[b[0] for b in SIZE_BUCKETS],
        ),
        "by_age": _group(
            objects,
            key=lambda o: _age_bucket(property_age(o.property_year, as_of.year)),
            order=[b[0] for b in AGE_BUCKETS] + ["год выпуска не указан"],
        ),
        "by_condition": _group(
            objects,
            key=lambda o: conditions[o.id],
            order=[CONDITION_NEW, CONDITION_USED, CONDITION_UNKNOWN],
        ),
        "by_brand": brand_top,
        "brands_recognized_share": _share(
            totals["sum"] - (not_recognized["sum"] if not_recognized else ZERO),
            totals["sum"],
        ),
        "by_branch": _group(
            objects, key=lambda o: o.branch_id, label=lambda o: o.branch
        ),
        "branch_by_insurer": _matrix(
            objects,
            lambda o: o.branch_id,
            lambda o: o.branch,
            lambda o: o.insurer_id,
            lambda o: o.insurer,
            columns_limit=MATRIX_INSURERS,
        ),
        "by_manager": managers[:TOP_MANAGERS],
        "managers_total": len(managers),
        "without_manager": _totals(without_manager),
        "ending_by_month": _ending_by_month(objects, as_of),
        "ending_soon_days": ENDING_SOON_DAYS,
        "by_end_year": end_years,
        "by_franchise": _group(
            objects,
            key=lambda o: _franchise_bucket(o.franchise),
            order=[b[0] for b in FRANCHISE_BUCKETS],
        ),
        "top_objects": [
            {
                "client": o.client,
                "description": o.description,
                "brand": brands[o.id],
                "insurer": o.insurer,
                "branch": o.branch,
                "sum": o.insured_sum,
                "end_date": o.end_date,
                "policy_number": o.policy_number,
            }
            for o in top_objects
        ],
    }


# --- Excel ---------------------------------------------------------------------


def build_property_report_workbook(report: Dict[str, Any]):
    """Книга Excel: лист на каждый блок отчёта, суммы — числами."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E9EEF5")
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(title, headers, rows, widths=None):
        ws = workbook.create_sheet(title[:31])
        ws.append([f"{title} · на {report['as_of']:%d.%m.%Y}"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(headers)
        for cell in ws[3]:
            cell.font, cell.fill = header_font, header_fill
        for row in rows:
            ws.append(row)
        for index, header in enumerate(headers):
            letter = ws.cell(row=3, column=index + 1).column_letter
            ws.column_dimensions[letter].width = (widths or {}).get(
                index, max(14, len(str(header)) + 2)
            )
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if isinstance(cell.value, Decimal):
                    cell.number_format = "#,##0" if abs(cell.value) >= 100 else "0.0"
                elif isinstance(cell.value, date):
                    cell.number_format = "DD.MM.YYYY"

    def group_rows(rows):
        return [
            [
                r["name"],
                r["count"],
                r["sum"].quantize(Decimal("1")),
                r["share"].quantize(Decimal("0.1")),
            ]
            for r in rows
        ]

    group_headers = ["", "Объектов", "Страховая стоимость, ₽", "Доля, %"]
    t = report["totals"]
    sheet(
        "Сводка",
        ["Показатель", "Значение"],
        [
            ["Объектов", t["count"]],
            ["Страховая стоимость, ₽", t["sum"].quantize(Decimal("1"))],
            ["Средняя стоимость объекта, ₽", t["avg"].quantize(Decimal("1"))],
            ["Лизингополучателей", t["lessees"]],
            ["Страховых компаний", t["insurers"]],
            [
                "Доля трёх крупнейших СК, %",
                report["concentration"]["top3_share"].quantize(Decimal("0.1")),
            ],
        ],
        {0: 34, 1: 22},
    )
    sheet(
        "Страховщики",
        [
            "Страховая компания",
            "Объектов",
            "Страховая стоимость, ₽",
            "Доля, %",
            "Средняя стоимость, ₽",
            f"Окончание в {report['ending_soon_days']} дн., ₽",
        ],
        [
            [
                r["name"],
                r["count"],
                r["sum"].quantize(Decimal("1")),
                r["share"].quantize(Decimal("0.1")),
                r["avg"].quantize(Decimal("1")),
                r["ending_soon"].quantize(Decimal("1")),
            ]
            for r in report["insurers"]
        ],
        {0: 26},
    )
    for title, key in (
        ("СК × вид", "insurer_by_type"),
        ("Филиал × СК", "branch_by_insurer"),
    ):
        matrix = report[key]
        sheet(
            title,
            [""] + matrix["headers"] + ["Итого"],
            [
                [row["name"]]
                + [c["sum"].quantize(Decimal("1")) for c in row["cells"]]
                + [row["total"].quantize(Decimal("1"))]
                for row in matrix["rows"]
            ],
            {0: 26},
        )
    for title, key in (
        ("Виды страхования", "by_type"),
        ("Размер объекта", "by_size"),
        ("Возраст", "by_age"),
        ("Новое и б-у", "by_condition"),
        ("Марки", "by_brand"),
        ("Филиалы", "by_branch"),
        ("Менеджеры", "by_manager"),
        ("Год окончания", "by_end_year"),
        ("Франшиза", "by_franchise"),
    ):
        sheet(title, group_headers, group_rows(report[key]), {0: 30})
    sheet(
        "Окончание по месяцам",
        ["Месяц", "Объектов", "Страховая стоимость, ₽"],
        [
            [r["label"], r["count"], r["sum"].quantize(Decimal("1"))]
            for r in report["ending_by_month"]
        ],
    )
    sheet(
        "Крупнейшие объекты",
        [
            "Лизингополучатель",
            "Имущество",
            "Марка",
            "Страховая компания",
            "Филиал",
            "Страховая стоимость, ₽",
            "Полис до",
        ],
        [
            [
                o["client"],
                o["description"],
                o["brand"],
                o["insurer"],
                o["branch"],
                o["sum"].quantize(Decimal("1")),
                o["end_date"],
            ]
            for o in report["top_objects"]
        ],
        {0: 30, 1: 50, 3: 20},
    )
    return workbook
