"""
ОСАГО в отчётах (этап 4), по карте влияния владельца.

С ОСАГО: отчёт по КВ, ежемесячный КВ, 3%, окончание срока, пролонгация,
конструктор (с фильтром). Без ОСАГО: четверговый, CSV, снимок по СС.
Экспорт полисов — ОСАГО только по галочке include_osago.
"""

import csv
import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.clients.models import Client
from apps.insurers.models import Branch, CommissionRate, InsuranceType, Insurer
from apps.policies.models import PaymentSchedule, Policy
from apps.reports.filters import PaymentExportFilter, PolicyExportFilter

pytestmark = [pytest.mark.integration, pytest.mark.django_db]

TODAY = date.today()
QUARTER = (TODAY.month - 1) // 3 + 1


@pytest.fixture
def staff_client(client, django_user_model):
    user = django_user_model.objects.create_user("staff", password="x", is_staff=True)
    client.force_login(user)
    return client


@pytest.fixture
def portfolio(db):
    osago_type = InsuranceType.objects.create(name="ОСАГО (тест)", is_osago=True)
    kasko_type = InsuranceType.objects.create(name="КАСКО (тест)")
    insurer = Insurer.objects.create(insurer_name="СК Отчёт")
    branch = Branch.objects.create(branch_name="Филиал Отчёт")
    client = Client.objects.create(client_name="ООО Отчёт", client_inn="7700000099")
    rates = {
        t: CommissionRate.objects.create(
            insurer=insurer, insurance_type=t, kv_percent=Decimal("10")
        )
        for t in (osago_type, kasko_type)
    }
    start = TODAY - timedelta(days=5)

    def make(number, ins_type, insurance_sum, end, **payment):
        policy = Policy.objects.create(
            policy_number=number,
            dfa_number=f"ДФА-{number}",
            client=client,
            insurer=insurer,
            branch=branch,
            insurance_type=ins_type,
            property_description=f"Имущество {number}",
            start_date=start,
            end_date=end,
            policy_uploaded=False,
        )
        PaymentSchedule.objects.create(
            policy=policy,
            year_number=1,
            installment_number=1,
            due_date=start,
            insurance_sum=insurance_sum,
            amount=Decimal("10000"),
            commission_rate=rates[ins_type],
            kv_rub=Decimal("1000"),
            **payment,
        )
        return policy

    end = start + timedelta(days=360)
    paid = {"paid_date": start}
    return {
        "insurer": insurer,
        "kasko": make("KASKO-REP", kasko_type, Decimal("2000000"), end, **paid),
        "osago": make("OSAGO-REP", osago_type, None, end, **paid),
        "end": end,
    }


def _approve_all():
    PaymentSchedule.objects.update(insurer_date=TODAY)


def xlsx_text(response):
    assert response.status_code == 200, response
    assert response["Content-Type"].startswith(
        "application/vnd.openxmlformats"
    ), response.get("Content-Type")
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    return " ".join(
        str(cell)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    )


class TestWithOsago:
    def test_three_percent_report(self, staff_client, portfolio):
        _approve_all()
        text = xlsx_text(
            staff_client.get(
                reverse("reports:export_three_percent_report"),
                {"three_percent_quarter": QUARTER, "three_percent_year": TODAY.year},
            )
        )
        assert "KASKO-REP" in text
        assert "OSAGO-REP" in text

    def test_monthly_kv_report(self, staff_client, portfolio):
        _approve_all()
        text = xlsx_text(
            staff_client.get(
                reverse("reports:export_monthly_kv_report"),
                {"kv_month": TODAY.month, "kv_year": TODAY.year},
            )
        )
        assert "OSAGO-REP" in text

    def test_commission_report(self, staff_client, portfolio):
        text = xlsx_text(
            staff_client.get(
                reverse("reports:export_commission_report"),
                {"insurer": portfolio["insurer"].pk},
            )
        )
        assert "KASKO-REP" in text
        assert "OSAGO-REP" in text

    def test_policy_expiration(self, staff_client, portfolio):
        text = xlsx_text(
            staff_client.get(
                reverse("reports:export_policy_expiration"),
                {
                    "date_from": (portfolio["end"] - timedelta(days=1)).isoformat(),
                    "date_to": (portfolio["end"] + timedelta(days=1)).isoformat(),
                },
            )
        )
        assert "OSAGO-REP" in text

    def test_prolongation(self, portfolio):
        from apps.billing import prolongation_services as prol

        end = portfolio["end"]
        policies = list(prol.get_prolongation_policies(end.year, end.month))
        assert portfolio["osago"] in policies

        batch = prol.get_or_create_batch(end.year, end.month)
        attachment = prol.build_prolongation_attachment(batch, policies)
        assert attachment


class TestWithoutOsago:
    def test_thursday_report(self, staff_client, portfolio):
        text = xlsx_text(staff_client.get(reverse("reports:export_thursday_report")))
        assert "KASKO-REP" in text
        assert "OSAGO-REP" not in text

    def test_policies_csv(self, staff_client, portfolio):
        start = TODAY - timedelta(days=30)
        response = staff_client.get(
            reverse("reports:export_policies_csv"),
            {"start_date": start.isoformat(), "end_date": TODAY.isoformat()},
        )
        assert response.status_code == 200
        content = response.content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(content), delimiter=";"))
        text = " ".join(" ".join(r) for r in rows)
        assert "KASKO-REP" in text
        assert "OSAGO-REP" not in text


class TestPolicyExportToggle:
    def test_osago_off_by_default(self, client, django_user_model, portfolio):
        client.force_login(django_user_model.objects.create_user("u", password="x"))
        text = xlsx_text(client.get(reverse("reports:export_policies")))
        assert "KASKO-REP" in text
        assert "OSAGO-REP" not in text

    def test_include_osago(self, client, django_user_model, portfolio):
        client.force_login(django_user_model.objects.create_user("u", password="x"))
        text = xlsx_text(
            client.get(reverse("reports:export_policies"), {"include_osago": "1"})
        )
        assert "OSAGO-REP" in text


class TestCustomExportFilter:
    @pytest.mark.parametrize(
        "value, expected",
        [
            ("", {"KASKO-REP", "OSAGO-REP"}),
            ("exclude", {"KASKO-REP"}),
            ("only", {"OSAGO-REP"}),
        ],
    )
    def test_policy_filter(self, portfolio, value, expected):
        qs = PolicyExportFilter({"osago": value}, queryset=Policy.objects.all()).qs
        assert set(qs.values_list("policy_number", flat=True)) == expected

    @pytest.mark.parametrize(
        "value, expected",
        [
            ("exclude", {"KASKO-REP"}),
            ("only", {"OSAGO-REP"}),
        ],
    )
    def test_payment_filter(self, portfolio, value, expected):
        qs = PaymentExportFilter(
            {"osago": value}, queryset=PaymentSchedule.objects.all()
        ).qs
        assert set(qs.values_list("policy__policy_number", flat=True)) == expected

    def test_custom_export_with_osago_payment(self, staff_client, portfolio):
        response = staff_client.post(
            reverse("reports:custom_export"),
            {
                "action": "export",
                "data_source": "payments",
                "fields": ["policy__policy_number", "insurance_sum", "amount"],
                "osago": "only",
            },
        )
        text = xlsx_text(response)
        assert "OSAGO-REP" in text
        assert "KASKO-REP" not in text
