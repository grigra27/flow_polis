"""
Четверговый отчёт: без ОСАГО и без закрытых ДФА (просьба учёта, 2026-10-16).

Обе части отчёта — «не подгружены документы» и «нет данных об оплате».
Раньше вторая часть отбирала неоплаченные взносы только по активности
полиса, и в отчёт попадали ОСАГО и полисы с закрытым ДФА.
"""

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.clients.models import Client
from apps.insurers.models import Branch, InsuranceType, Insurer
from apps.policies.models import PaymentSchedule, Policy

pytestmark = [pytest.mark.integration, pytest.mark.django_db]

TODAY = date.today()


@pytest.fixture
def staff_client(client, django_user_model):
    client.force_login(
        django_user_model.objects.create_user("staff", password="x", is_staff=True)
    )
    return client


@pytest.fixture
def make_policy(db):
    kasko = InsuranceType.objects.create(name="КАСКО (четверг)")
    osago = InsuranceType.objects.create(name="ОСАГО (четверг)", is_osago=True)
    insurer = Insurer.objects.create(insurer_name="СК Четверг")
    branch = Branch.objects.create(branch_name="Филиал Четверг")
    client = Client.objects.create(client_name="ООО Четверг", client_inn="7700000077")
    start = TODAY - timedelta(days=40)

    def make(number, *, is_osago=False, uploaded=True, paid=False, dfa_active=True):
        policy = Policy.objects.create(
            policy_number=number,
            dfa_number=f"ДФА-{number}",
            client=client,
            insurer=insurer,
            branch=branch,
            insurance_type=osago if is_osago else kasko,
            property_description=f"Имущество {number}",
            start_date=start,
            end_date=start + timedelta(days=364),
            policy_uploaded=uploaded,
            dfa_active=dfa_active,
            dfa_deactivation_date=None if dfa_active else TODAY - timedelta(days=10),
        )
        PaymentSchedule.objects.create(
            policy=policy,
            year_number=1,
            installment_number=1,
            due_date=start,
            insurance_sum=None if is_osago else Decimal("1000000"),
            amount=Decimal("10000"),
            kv_rub=Decimal("1000"),
            paid_date=start if paid else None,
        )
        return policy

    return make


def report_policy_numbers(response):
    assert response.status_code == 200, response
    ws = load_workbook(io.BytesIO(response.content), read_only=True).active
    return {row[0] for row in ws.iter_rows(min_row=4, values_only=True) if row[0]}


def test_unpaid_section_skips_osago_and_closed_dfa(staff_client, make_policy):
    make_policy("UNPAID-OK")
    make_policy("UNPAID-OSAGO", is_osago=True)
    make_policy("UNPAID-DFA-CLOSED", dfa_active=False)

    numbers = report_policy_numbers(
        staff_client.get(reverse("reports:export_thursday_report"))
    )

    assert "UNPAID-OK" in numbers
    assert "UNPAID-OSAGO" not in numbers
    assert "UNPAID-DFA-CLOSED" not in numbers


def test_not_uploaded_section_skips_osago_and_closed_dfa(staff_client, make_policy):
    make_policy("DOCS-OK", uploaded=False, paid=True)
    make_policy("DOCS-OSAGO", is_osago=True, uploaded=False, paid=True)
    make_policy("DOCS-DFA-CLOSED", uploaded=False, paid=True, dfa_active=False)

    numbers = report_policy_numbers(
        staff_client.get(reverse("reports:export_thursday_report"))
    )

    assert "DOCS-OK" in numbers
    assert "DOCS-OSAGO" not in numbers
    assert "DOCS-DFA-CLOSED" not in numbers


def test_future_cutoff_date_keeps_the_same_rules(staff_client, make_policy):
    """Дата среза в будущем (как в выгрузке учёта: «дата среза 25.10»)."""
    make_policy("CUTOFF-OK")
    make_policy("CUTOFF-OSAGO", is_osago=True)
    make_policy("CUTOFF-DFA-CLOSED", dfa_active=False)

    numbers = report_policy_numbers(
        staff_client.get(
            reverse("reports:export_thursday_report"),
            {"payment_date": (TODAY + timedelta(days=9)).isoformat()},
        )
    )

    assert numbers & {"CUTOFF-OK", "CUTOFF-OSAGO", "CUTOFF-DFA-CLOSED"} == {"CUTOFF-OK"}
